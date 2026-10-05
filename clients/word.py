import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from pygments import lex
from pygments.lexers import PythonLexer
from pygments.styles import get_style_by_name

PAGE_WIDTH_CM = 15.25

CODE_STYLE = "Code"
CODE_FONT = "Courier New"  # есть на любом компьютере, в отличие от JetBrains Mono
CODE_FONT_SIZE = Pt(10)
PYGMENTS_STYLE = get_style_by_name("friendly")


class WordClient:
    def __init__(self, doc_path: str):
        self.doc = Document(docx=doc_path)
        self._ensure_code_style()

    def _ensure_code_style(self) -> None:
        if CODE_STYLE in [s.name for s in self.doc.styles]:
            return
        style = self.doc.styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = self.doc.styles["Normal"]
        style.font.name = CODE_FONT
        style.font.size = CODE_FONT_SIZE
        pf = style.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(0)
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    def add_heading1(self, title: str) -> None:
        self.doc.add_heading(title, level=1)
        self.doc.add_paragraph()

    def add_heading2(self, title: str, heading_no: int) -> None:
        self.doc.add_heading(
            f"2.{heading_no}. Решения задач на тему «{title}»",
            level=2,
        )
        self.doc.add_paragraph()

    def _add_problem(self, title: str, descr: str) -> None:
        self.doc.add_paragraph(f"«{title}».")
        self.doc.add_paragraph(f"{descr}")
        self.doc.add_paragraph()

    def add_code(self, code: str) -> None:
        """Вставляет код текстом с подсветкой синтаксиса (один абзац стиля Code)."""
        p = self.doc.add_paragraph(style=CODE_STYLE)
        for token_type, value in lex(code.strip("\n"), PythonLexer()):
            token_style = PYGMENTS_STYLE.style_for_token(token_type)
            run = p.add_run(value)
            if token_style["color"]:
                run.font.color.rgb = RGBColor.from_string(token_style["color"])
            run.bold = token_style["bold"]
            run.italic = token_style["italic"]

    def add_solution(self, no: int, title: str, descr: str, code: str) -> None:
        """Решение в виде листинга: код текстом, его можно копировать."""
        self._add_problem(title, descr)

        label = self.doc.add_paragraph(f"Листинг 2.{no} – Решение задачи «{title}»")
        label.paragraph_format.keep_with_next = True
        self.add_code(code)

        self.doc.add_paragraph()

    def add_solution_picture(
        self, no: int, title: str, descr: str, img_path: str
    ) -> None:
        """Решение в виде картинки (старый режим, флаг --images)."""
        self._add_problem(title, descr)

        pic_p = self.doc.add_paragraph()
        pic_p.add_run().add_picture(img_path, width=Cm(PAGE_WIDTH_CM))
        pic_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        label = self.doc.add_paragraph(f"Рисунок 2.{no} – Решение задачи «{title}»")
        label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        self.doc.add_paragraph()

    def add_page_break(self):
        self.doc.add_page_break()

    def save(self, doc_name: str) -> str:
        os.makedirs("my_docs", exist_ok=True)
        if not doc_name.endswith(".docx"):
            doc_name += ".docx"
        path = f"my_docs/{doc_name}"
        self.doc.save(path)
        return path
