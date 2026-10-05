"""Вставка кода в .docx текстом с подсветкой синтаксиса: python-docx + Pygments."""

from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from pygments import lex
from pygments.lexers import get_lexer_by_name
from pygments.styles import get_style_by_name

CODE_STYLE = "Code"

# Порядок детей <w:pPr> строгий: Word ругается на «нечитаемое содержимое»,
# если, например, <w:shd> окажется после <w:spacing>.
_PPR_TAIL = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _set_all_fonts(rpr, name: str) -> None:
    """style.font.name задает только ascii/hAnsi – добиваем cs и eastAsia."""
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _add_box(ppr, fill: str = "F6F8FA", border: str = "D0D7DE") -> None:
    """Рамка и заливка абзаца – чтобы листинг выглядел блоком."""
    pbdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")  # в восьмых пункта: 4 = 0,5 pt
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), border)
        pbdr.append(el)
    ppr.insert_element_before(pbdr, "w:shd", *_PPR_TAIL)

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.insert_element_before(shd, *_PPR_TAIL)


def ensure_code_style(doc: DocumentObject, font: str = "Courier New", size: int = 10) -> None:
    if CODE_STYLE in [s.name for s in doc.styles]:
        return
    style = doc.styles.add_style(CODE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(size)
    style.font.no_proof = True  # без красных волнистых подчеркиваний
    _set_all_fonts(style.element.get_or_add_rPr(), font)

    pf = style.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = Pt(0)
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT  # в шаблоне Normal может быть «по ширине»
    _add_box(style.element.get_or_add_pPr())


def _runs(code: str, language: str, style_name: str):
    """Токены Pygments -> склеенные фрагменты (текст, цвет, bold, italic)."""
    style = get_style_by_name(style_name)
    chunks: list[list] = []
    for token_type, value in lex(code, get_lexer_by_name(language, ensurenl=False)):
        s = style.style_for_token(token_type)
        fmt = (s["color"], s["bold"], s["italic"])
        if chunks and (chunks[-1][1] == fmt or value.isspace()):
            chunks[-1][0] += value  # пробелы и одинаковый стиль – в тот же run
        else:
            chunks.append([value, fmt])
    return chunks


def add_code(
    doc: DocumentObject,
    code: str,
    language: str = "python",
    style_name: str = "friendly",
) -> Paragraph:
    ensure_code_style(doc)
    p = doc.add_paragraph(style=CODE_STYLE)
    for text, (color, bold, italic) in _runs(code.strip("\n"), language, style_name):
        run = p.add_run(text)  # \n -> <w:br/>, \t -> <w:tab/>
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        run.bold = bold
        run.italic = italic
    return p


def _add_field(paragraph: Paragraph, instr: str, cached: str) -> None:
    """Поле Word (как «Вставить название»): begin – instrText – separate – результат – end."""
    def fld(kind: str):
        r = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)

    fld("begin")
    r = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    r._r.append(instr_el)
    fld("separate")
    paragraph.add_run(cached)  # то, что видно до обновления полей
    fld("end")


def add_listing(
    doc: DocumentObject,
    code: str,
    title: str,
    number: int,
    language: str = "python",
) -> None:
    """Подпись «Листинг N – …» с автонумерацией SEQ + сам код."""
    caption = doc.add_paragraph("Листинг ")
    _add_field(caption, "SEQ Листинг \\* ARABIC", str(number))
    caption.add_run(f" – {title}")
    caption.paragraph_format.keep_with_next = True

    add_code(doc, code, language)
