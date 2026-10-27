"""Запись отчета в docx: заголовки уроков, условия, листинги или картинки кода.

Нумерация и подписи – шаблоны строк, поэтому под другой формат отчета ничего переписывать
не нужно: Report(part_no=1, lesson_heading="{part}.{n} Решения задач на тему «{title}»").
"""

import os

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from .docx_code import add_code, ensure_code_style
from .models import CodeSolution

PAGE_WIDTH_CM = 15.25


class Report:
    def __init__(
        self,
        template: str | None = None,
        part_no: int = 2,
        lesson_heading: str = "{part}.{n}. Решения задач на тему «{title}»",
        listing_caption: str = "Листинг {part}.{n} – Решение задачи «{title}»",
        picture_caption: str = "Рисунок {part}.{n} – Решение задачи «{title}»",
    ):
        """template – свой docx с титульником, введением и стилями; None – пустой документ."""
        self.doc = Document(template)
        ensure_code_style(self.doc)
        self.part_no = part_no
        self.lesson_heading = lesson_heading
        self.listing_caption = listing_caption
        self.picture_caption = picture_caption
        self.lessons = 0
        self.solutions = 0

    def add_heading(self, title: str, level: int = 1) -> None:
        self.doc.add_heading(title, level=level)
        self.doc.add_paragraph()

    def add_paragraph(self, text: str = "") -> Paragraph:
        return self.doc.add_paragraph(text)

    def add_lesson(self, title: str) -> None:
        self.lessons += 1
        heading = self.lesson_heading.format(part=self.part_no, n=self.lessons, title=title)
        self.add_heading(heading, level=2)

    def _add_problem(self, title: str, descr: str) -> None:
        self.doc.add_paragraph(f"«{title}».")
        if descr:
            self.doc.add_paragraph(descr)
        self.doc.add_paragraph()

    def _caption(self, template: str, title: str) -> str:
        return template.format(part=self.part_no, n=self.solutions, title=title)

    def add_solution(self, solution: CodeSolution) -> None:
        """Решение в виде листинга: код текстом, его можно копировать."""
        self.solutions += 1
        self._add_problem(solution.title, solution.description)

        label = self.doc.add_paragraph(self._caption(self.listing_caption, solution.title))
        label.paragraph_format.keep_with_next = True
        add_code(self.doc, solution.code, solution.language)

        self.doc.add_paragraph()

    def add_solution_picture(self, solution: CodeSolution) -> None:
        """Решение в виде картинки (старый режим, флаг --images)."""
        self.solutions += 1
        self._add_problem(solution.title, solution.description)

        pic_p = self.doc.add_paragraph()
        pic_p.add_run().add_picture(solution.img_path, width=Cm(PAGE_WIDTH_CM))
        pic_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        label = self.doc.add_paragraph(self._caption(self.picture_caption, solution.title))
        label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        self.doc.add_paragraph()

    def add_page_break(self) -> None:
        self.doc.add_page_break()

    def save(self, path: str) -> str:
        if not path.endswith(".docx"):
            path += ".docx"
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)
        self.doc.save(path)
        return path
