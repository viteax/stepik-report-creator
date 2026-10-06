from docx import Document
from pygments import lex
from pygments.lexers import get_lexer_by_name

from docx_code import CODE_STYLE, _runs, add_code, add_listing, ensure_code_style


def test_code_style_is_created_once():
    doc = Document()
    ensure_code_style(doc)
    ensure_code_style(doc)

    assert [s.name for s in doc.styles].count(CODE_STYLE) == 1


def test_code_is_one_paragraph_with_all_text():
    doc = Document()
    code = "def solve(s):\n    stack = []\n    return not stack\n"

    p = add_code(doc, code)

    assert p.style.name == CODE_STYLE
    assert "stack = []" in p.text
    assert p.text.count("\n") == 2


def test_runs_are_merged():
    code = "def f(x):\n    return x + 1\n"

    chunks = _runs(code, "python", "friendly")

    raw_tokens = list(lex(code, get_lexer_by_name("python", ensurenl=False)))
    assert len(chunks) < len(raw_tokens)  # склеено, а не по токену
    assert "".join(text for text, _ in chunks) == code.strip("\n")


def test_listing_has_seq_field_and_keeps_with_next():
    doc = Document()

    add_listing(doc, "print(1)", title="Вывод", number=3)

    caption = doc.paragraphs[0]
    assert caption.text.startswith("Листинг 3")
    assert caption.paragraph_format.keep_with_next
    assert "SEQ Листинг" in caption._p.xml
