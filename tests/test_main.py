import os
import sys
from types import SimpleNamespace

import pytest
from docx import Document

from stepik_report import cli
from stepik_report.models import Reply

CODE = "def f(x):\n    return x * 2\n\nprint(f(int(input())))\n"


class FakeStepik:
    """Подставной клиент: те же методы, что у StepikClient, но без сети."""

    def get_section_id(self, course_id, section_no):
        return 1

    def get_section(self, id):
        return SimpleNamespace(title="Графы и деревья", units=[1, 2])

    def get_section_lessons(self, section):
        return [
            SimpleNamespace(title=f"Урок {i}", steps=[i * 10, i * 10 + 1])
            for i in range(1, 3)
        ]

    def get_steps(self, ids):
        return [
            SimpleNamespace(
                id=i,
                block=SimpleNamespace(
                    name="code",
                    text=f"<h2>Задача {i}</h2><p>Условие {i}</p><p>Формат входных данных</p>",
                ),
            )
            for i in ids
        ]

    def get_solution(self, step_id):
        return Reply(code=CODE, language="python3")


@pytest.fixture
def workdir(tmp_path, monkeypatch):
    os.symlink(os.path.abspath("assets"), tmp_path / "assets")
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(cli, "make_client", FakeStepik)
    return tmp_path


def test_main_builds_report_with_code_as_text(workdir, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["main.py", "7"])

    cli.main()

    doc = Document(str(workdir / "my_docs" / "7-Графы-и-деревья.docx"))
    code = [p.text for p in doc.paragraphs if p.style.name == "Code"]
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Листинг")]
    assert len(code) == 4
    assert all("return x * 2" in c for c in code)
    assert [c.split(" –")[0] for c in captions] == [f"Листинг 2.{i}" for i in range(1, 5)]
    assert "Задач в отчете: 4" in capsys.readouterr().out


def test_main_images_mode(workdir, monkeypatch, capsys):
    monkeypatch.setattr(sys, "argv", ["main.py", "7", "--images"])

    cli.main()

    doc = Document(str(workdir / "my_docs" / "7-Графы-и-деревья.docx"))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Рисунок 2.")]  # «1.N» – из шаблона
    assert len(captions) == 4
    assert "Задач в отчете: 4" in capsys.readouterr().out
